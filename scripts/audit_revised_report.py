import docx

doc = docx.Document("docs/ARGUS_Comprehensive_Technical_Report_REVISED.docx")
full_text = []
for p in doc.paragraphs:
    full_text.append(p.text)
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            full_text.append(c.text)

combined = " ".join(full_text).lower()

forbidden = [
    "enterprise-grade",
    "calibrated probability",
    "calibrated linear baseline",
    "density-based isolation forest",
    "500,000 fast path",
    "automatically approve",
    "counterparty accounts temporarily frozen",
    "real-world operational efficiency",
    "production-ready",
    "implemented jwt",
    "implemented hmac",
    "implemented postgresql",
    "implemented esp32",
    "18 → 12 → 6 → 12 → 18",
    "18-12-6-12-18",
]

print("=== FORBIDDEN PHRASES AUDIT ===")
found_forbidden = False
for phrase in forbidden:
    count = combined.count(phrase.lower())
    if count > 0:
        print(f'FAIL: Found {count} occurrences of forbidden phrase: "{phrase}"')
        found_forbidden = True
    else:
        print(f'PASS: 0 occurrences of "{phrase.encode("ascii", "replace").decode()}"')

if not found_forbidden:
    print(">>> ALL FORBIDDEN PHRASES AUDIT PASSED! <<<")

words = sum(len(p.text.split()) for p in doc.paragraphs) + sum(len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells)
shapes = [p for p in doc.paragraphs if "graphicData" in p._element.xml]

print("\n=== DOCUMENT METRICS ===")
print(f"Total Words (text + tables): {words}")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Embedded Figures: {len(shapes)}")

# Check occurrence of actual architecture
arch_count = combined.count("18 → 64 → 32 → 16 → 32 → 64 → 18") + combined.count("18-64-32-16-32-64-18")
print(f"Verified Autoencoder Architecture occurrences: {arch_count}")
