#!/usr/bin/env python3
"""
A.R.G.U.S. — Comprehensive Technical Engineering & Progress Report Generator (.docx)
Technical and Academic Correction Pass (Milestones 0 to 4).
Generates:
  - docs/ARGUS_Comprehensive_Technical_Report_REVISED.docx
  - docs/ARGUS_Comprehensive_Technical_Report.docx
"""

import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Palette
NAVY = RGBColor(26, 54, 93)       # #1A365D - Heading 1 & Main Title
BLUE = RGBColor(43, 108, 176)     # #2B6CB0 - Heading 2
TEAL = RGBColor(49, 151, 149)     # #319795 - Heading 3
CHARCOAL = RGBColor(45, 55, 72)   # #2D3748 - Body
MUTED = RGBColor(113, 128, 150)   # #718096 - Captions & Subtitles
WHITE = RGBColor(255, 255, 255)

HEX_NAVY = "1A365D"
HEX_LIGHT_BG = "F7FAFC"
HEX_BORDER = "CBD5E0"
HEX_CALLOUT_BG = "EDF2F7"
HEX_CALLOUT_BORDER = "3182CE"
HEX_DARK_BOX = "1A202C"


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding (in twentieths of a point / dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_callout(doc, text, title="CRITICAL NOTICE"):
    """Adds a callout box with a colored left border and subtle background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = BLUE
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = CHARCOAL
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)


def style_table(table, col_widths, header_bg=HEX_NAVY):
    """Applies modern academic styling to tables with alternating row shading."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = WHITE
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = HEX_LIGHT_BG if (r_idx % 2 == 1) else "FFFFFF"
        for i, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[i])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = CHARCOAL


def build_report():
    doc = Document()

    # Configure Margins (1 inch everywhere)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("A.R.G.U.S. — Technical Progress Report | B.Tech AIML (2026–27)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Automated Risk Assessment & Anomaly Detection System — Academic Progress Report (Milestones 0–4)")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED

    # --------------------------------------------------------------------------
    # COVER PAGE
    # --------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("A.R.G.U.S.")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(32)
    r_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Automated Risk Assessment & Anomaly Detection System")
    r_sub.bold = True
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = BLUE

    p_tagline = doc.add_paragraph()
    p_tagline.paragraph_format.space_before = Pt(0)
    p_tagline.paragraph_format.space_after = Pt(28)
    p_tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tagline = p_tagline.add_run("AI-Based Real-Time Fraud Detection and Transaction Risk Monitoring System\nComprehensive Technical Engineering & Progress Report (Milestones 0 – 4)")
    r_tagline.italic = True
    r_tagline.font.name = "Calibri"
    r_tagline.font.size = Pt(11.5)
    r_tagline.font.color.rgb = MUTED

    # Metadata Table
    meta_table = doc.add_table(rows=7, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [2.2, 4.3]
    
    meta_data = [
        ("Project Name", "A.R.G.U.S. (Automated Risk Assessment & Anomaly Detection System)"),
        ("Academic Programme", "B.Tech Artificial Intelligence & Machine Learning"),
        ("Class & Academic Year", "B.Tech AIML – C | Academic Year 2026–27 (Odd Semester)"),
        ("Engineering Team", "Savar Shetty, Pranav Pawar, Siddhesh Gire, Devraj Misal, Atharva Morbale"),
        ("Faculty Mentor / Guide", "Prof. Sunil Kale"),
        ("Completed Milestones", "Milestone 0 → Milestone 1 → Milestone 2 → Milestone 3 → Milestone 3A → Milestone 4"),
        ("Report Classification", "Technical Progress Report: Architecture, ML/DL Training, Audit & Risk Engine"),
    ]
    
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(col_w[0])
        c1.width = Inches(col_w[1])
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "F7FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = NAVY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = CHARCOAL

    doc.add_page_break()

    # --------------------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & PROJECT CHARTER
    # --------------------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary & Project Charter", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "A.R.G.U.S. (Automated Risk Assessment & Anomaly Detection System) is a modular AI-based fraud detection and "
        "transaction risk monitoring system engineered for financial services, mobile money platforms, and point-of-sale "
        "payment gateways. In modern digital finance, static rule engines struggle with binary cutoffs that either generate "
        "excessive false positives or miss novel fraudulent exfiltration patterns. A.R.G.U.S. addresses this by synthesizing "
        "supervised classification, tree-based unsupervised anomaly detection, and deep autoencoder reconstruction error into "
        "a normalized, continuous Risk Score (0 to 100) paired with an operational tri-state decision policy: APPROVE, REVIEW, "
        "and BLOCK."
    )

    doc.add_paragraph(
        "This document serves as the official Technical Progress Report for Milestones 0 through 4, detailing the engineering "
        "progress completed during development:\n"
        "• Milestone 0 (Environment Initialization): Established a reproducible Python 3.13.7 virtual environment, verified "
        "compatibility across PyTorch, XGBoost, Scikit-learn, and FastAPI, and configured strict Git repository hygiene.\n"
        "• Milestone 1 (PaySim Acquisition & EDA): Performed empirical exploratory data analysis on the 6.36-million row PaySim "
        "synthetic dataset, documenting 100% fraud concentration in TRANSFER and CASH_OUT, 97.55% sender account liquidation, "
        "and a 773.7:1 extreme class imbalance.\n"
        "• Milestone 2 (Feature Engineering & Splitting): Engineered an 18-feature domain pipeline, isolated targets to eliminate "
        "leakage, established a chronologically partitioned Train (Steps 1–520), Validation (Steps 521–631), and Test (Steps 632–743) "
        "split, and fitted the feature scaler strictly on training data.\n"
        "• Milestone 3 (Model Preparation & Training): Implemented and trained four complementary models: Logistic Regression, "
        "XGBoost, Isolation Forest, and a PyTorch Deep Autoencoder (18 → 64 → 32 → 16 → 32 → 64 → 18).\n"
        "• Milestone 3A (Forensic Model Audit): Investigated the near-perfect XGBoost metrics, attributed 92.25% of feature gain to "
        "PaySim's deterministic liquidation rules, audited the singular test false negative (row 3583), and resolved a PyTorch "
        "tensor writability warning without retraining.\n"
        "• Milestone 4 (Risk Scoring & Decision Engine): Implemented a weighted multi-signal aggregator (0.50 XGBoost, 0.20 Autoencoder, "
        "0.15 Isolation Forest, 0.15 Logistic Regression), producing a 0–100 score that captured 99.92% of fraud cases with zero "
        "legitimate transactions blocked across 78,701 validation transactions."
    )

    add_callout(
        doc,
        "MANDATORY ARCHITECTURAL SCOPE & INTEGRITY AUDIT:\n"
        "This progress report strictly distinguishes between what has been IMPLEMENTED, what has been ANALYZED / DESIGNED, "
        "and what is PLANNED / FUTURE:\n"
        "• IMPLEMENTED (Milestones 0–4): Environment setup, PaySim EDA, 18-feature pipeline, chronological dataset split, "
        "four trained ML/DL models, forensic audit, multi-signal Risk Engine, signal normalization, tri-state policy, and 20 unit tests.\n"
        "• ANALYZED / DESIGNED (Planned for Milestones 5–7): Relational 3NF PostgreSQL database schema, FastAPI REST API architecture, "
        "ESP32 edge IoT POS architecture, security mechanisms (JWT, HMAC-SHA256, TLS, RBAC), and OOP domain models.\n"
        "• PLANNED / FUTURE (Milestones 8–9): Live database migration, production REST service deployment, web analyst dashboard, "
        "hardware flashing of ESP32 microcontrollers, and continuous learning pipeline.",
        "ACADEMIC INTEGRITY & IMPLEMENTATION BOUNDARY"
    )

    h2 = doc.add_heading("Academic Subject Integration Framework", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "A.R.G.U.S. directly integrates the six core subjects of the B.Tech Artificial Intelligence and Machine Learning curriculum, "
        "demonstrating rigorous cross-disciplinary engineering:"
    )

    # Subject Integration Table
    subj_table = doc.add_table(rows=7, cols=3)
    s_col_w = [1.8, 1.8, 2.9]
    subj_data = [
        ("Curricular Discipline", "Core Academic Concepts", "Concrete A.R.G.U.S. Implementation & Architecture"),
        ("1. Deep Learning", "Autoencoders, encoder/decoder networks, latent representations, reconstruction error, unsupervised anomaly detection.", "[IMPLEMENTED] PyTorch Deep Autoencoder (18 → 64 → 32 → 16 → 32 → 64 → 18) with LeakyReLU and BatchNorm1d, trained on normal transactions to identify anomalies via MSE reconstruction loss."),
        ("2. DBMS", "Relational modelling, 3NF normalization, primary/foreign keys, ACID transactions, persistence, audit logging.", "[DESIGNED / PLANNED — Milestone 5] 3NF schema designed across 11 entities (users, transactions, devices, merchants, risk_assessments, audit_logs, etc.) for PostgreSQL integration."),
        ("3. Computational Techniques", "Feature transformations, logarithmic compression, scaling/normalization, cyclical mathematical encoding, statistical analysis, weighted aggregation.", "[IMPLEMENTED] Logarithmic amount compression, Z-score standardization, cyclical trigonometric encoding (sin/cos of hour), normal error distribution modeling (mean, std, P95), and weighted linear risk aggregation."),
        ("4. Information Security", "Authentication, authorization, RBAC, API validation, password hashing, HMAC payload integrity, secure communications, audit trails.", "[IMPLEMENTED: M0–M4] Environment secret isolation (.env.example), .gitignore shielding, deterministic seeds. [DESIGNED / PLANNED: M5–M7] JWT bearer auth, RBAC, HMAC-SHA256 POS signing, TLS 1.3, SHA-256 audit chaining."),
        ("5. IoT (Internet of Things)", "Hardware telemetry, edge point-of-sale terminals, tamper sensors, secure payload transmission, RTC timestamps.", "[DESIGNED / PLANNED — Milestone 7] Planned ESP32/POS integration with hardware tamper switch, device_id, merchant_id, RTC timestamps, and HMAC payload signing. Physical ESP32 integration is pending."),
        ("6. OOPs (Object-Oriented Programming)", "Encapsulation, domain modeling, service abstractions, modular class hierarchies, separation of concerns.", "[IMPLEMENTED] TransactionRiskEngine, ModelSignalNormalizer, RiskAssessment, DeepAutoencoder, RiskEngineConfig. [DESIGNED / PLANNED] Domain entities (User, Transaction, Device) and services (TransactionService, FraudDetectionService)."),
    ]
    for idx, r_data in enumerate(subj_data):
        row = subj_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(subj_table, s_col_w)

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 2. SYSTEM ARCHITECTURE & WORKFLOW
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("2. System Architecture & Component Flow", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "A.R.G.U.S. is structured as a modular processing pipeline. Incoming transactions pass through validation, feature "
        "engineering, parallel model inference, and weighted risk aggregation before producing an operational decision. "
        "Components currently implemented versus planned for future milestones are clearly demarcated below:"
    )

    # Visual Architecture Box
    arch_box = doc.add_table(rows=1, cols=1)
    arch_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_arch = arch_box.cell(0, 0)
    c_arch.width = Inches(6.5)
    set_cell_background(c_arch, HEX_DARK_BOX)
    set_cell_margins(c_arch, top=140, bottom=140, left=160, right=160)
    
    p_diag = c_arch.paragraphs[0]
    diag_text = (
        "                    ┌────────────────────────────────────────────────┐\n"
        "                    │               TRANSACTION SOURCES              │\n"
        "                    │   Web / Mobile / Payment Gateway / IoT POS     │ [PLANNED / M7]\n"
        "                    └───────────────────────┬────────────────────────┘\n"
        "                                            │ HTTPS / TLS [PLANNED / M6]\n"
        "                                            ▼\n"
        "                    ┌────────────────────────────────────────────────┐\n"
        "                    │              FASTAPI / API LAYER               │\n"
        "                    │      Pydantic Validation & Security Gateway    │ [DESIGNED / M6]\n"
        "                    └───────────────────────┬────────────────────────┘\n"
        "                                            │ Validated Payload\n"
        "                                            ▼\n"
        "                    ┌────────────────────────────────────────────────┐\n"
        "                    │              FEATURE ENGINEERING               │\n"
        "                    │    18 Features & Behavioral Domain Signals     │ [IMPLEMENTED / M2]\n"
        "                    └───────────────────────┬────────────────────────┘\n"
        "                                            │\n"
        "                 ┌──────────────────────────┼──────────────────────────┐\n"
        "                 ▼                          ▼                          ▼\n"
        "        ┌─────────────────┐        ┌─────────────────┐        ┌─────────────────┐\n"
        "        │     XGBOOST     │        │ ISOLATION FOREST│        │ DEEP AUTOENCODER│\n"
        "        │   CLASSIFIER    │        │  Tree Isolation │        │ Reconstruction  │\n"
        "        │  Weight: 0.50   │        │  Weight: 0.15   │        │  Weight: 0.20   │\n"
        "        │  [IMPLEMENTED]  │        │  [IMPLEMENTED]  │        │  [IMPLEMENTED]  │\n"
        "        └────────┬────────┘        └────────┬────────┘        └────────┬────────┘\n"
        "                 │                          │                          │\n"
        "                 └──────────────────────────┼──────────────────────────┘\n"
        "                                            │   ┌─────────────────┐\n"
        "                                            │   │    LOGISTIC     │\n"
        "                                            ├───►   REGRESSION    │\n"
        "                                            │   │  Weight: 0.15   │\n"
        "                                            │   │  [IMPLEMENTED]  │\n"
        "                                            │   └─────────────────┘\n"
        "                                            ▼\n"
        "                    ┌────────────────────────────────────────────────┐\n"
        "                    │                  RISK ENGINE                   │\n"
        "                    │   Signal Normalization & Weighted Aggregation  │ [IMPLEMENTED / M4]\n"
        "                    └───────────────────────┬────────────────────────┘\n"
        "                                            │\n"
        "                                            ▼\n"
        "                    ┌────────────────────────────────────────────────┐\n"
        "                    │               RISK SCORE (0–100)               │ [IMPLEMENTED / M4]\n"
        "                    └───────────────────────┬────────────────────────┘\n"
        "                                            │\n"
        "                 ┌──────────────────────────┼──────────────────────────┐\n"
        "                 ▼                          ▼                          ▼\n"
        "              APPROVE                     REVIEW                     BLOCK\n"
        "           (Score < 40)             (40 <= Score < 70)           (Score >= 70)\n"
        "       Immediate Settlement        Analyst Investigation       Policy Rejection\n"
    )
    r_diag = p_diag.add_run(diag_text)
    r_diag.font.name = "Consolas"
    r_diag.font.size = Pt(7.8)
    r_diag.font.color.rgb = RGBColor(226, 232, 240)

    doc.add_paragraph()

    doc.add_paragraph(
        "Current End-to-End Processing Flow:\n"
        "1. Ingress & Routing: A transaction payload containing transaction type, monetary amount, sender/recipient identifiers, "
        "and account balances enters the processing pipeline.\n"
        "2. Handling Non-Modeled Types: Because the current supervised fraud-modeling subspace contains TRANSFER and CASH_OUT "
        "transactions, the remaining transaction types (PAYMENT, CASH_IN, DEBIT) will be handled through a dedicated low-latency "
        "business-rule path. The exact rules, safeguards, and operational policy will be finalized during API integration. [DESIGNED / PLANNED]\n"
        "3. Feature Extraction: For TRANSFER and CASH_OUT transactions, the pipeline computes 18 engineered features, transforms "
        "monetary values via natural logarithm, derives cyclical temporal encodings, and calculates ledger balance discrepancies. [IMPLEMENTED]\n"
        "4. Multi-Model Inference: The 18 features are evaluated concurrently by XGBoost, Logistic Regression, Isolation Forest, and Deep Autoencoder. [IMPLEMENTED]\n"
        "5. Normalization: Raw inference outputs (probabilities, decision margins, reconstruction MSE) are mapped onto a standard [0.0, 1.0] scale. [IMPLEMENTED]\n"
        "6. Risk Aggregation: Signals are combined via a linear weighted sum and multiplied by 100 to yield a continuous Risk Score. [IMPLEMENTED]\n"
        "7. Decision & Explainability: The score triggers an automated tri-state policy (Approve / Review / Block), paired with "
        "human-readable risk factors and observable evidence. [IMPLEMENTED]"
    )

    # --------------------------------------------------------------------------
    # 3. MILESTONE 0: ENVIRONMENT & REPOSITORY SETUP
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("3. Milestone 0: Environment & Repository Setup", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "Reproducibility, dependency isolation, and version-controlled hygiene form the foundation of scientific software engineering. "
        "In Milestone 0, the project workspace was initialized from an empty state into a disciplined, version-controlled architecture."
    )

    doc.add_paragraph(
        "Host Environment & Toolchain Specifications:\n"
        "• Operating System: Microsoft Windows 10/11 (64-bit architecture), PowerShell Core.\n"
        "• Python Runtime: Python 3.13.7 (64-bit executable) in an isolated virtual environment (.venv/).\n"
        "• Version Control: Git 2.55.0.windows.2, repository initialized at commit 0f1b9c7."
    )

    doc.add_paragraph(
        "Dependency Verification & Ecosystem Compatibility:\n"
        "Prior to constructing the pipeline, an automated dependency verification confirmed that official pre-compiled CPython 3.13 "
        "x86_64 binary wheels exist on PyPI for all required libraries, avoiding local compilation overhead:\n"
        "• PyTorch (torch 2.14.0+cpu): Pre-compiled binary wheel installed cleanly; tensor operations and autograd verified.\n"
        "• XGBoost (xgboost 3.4.1): Native C++ gradient boosting engine with OpenMP multi-threading verified on 64-bit Windows.\n"
        "• Scikit-learn (scikit-learn 1.9.1): Pre-compiled Cython extensions verified for IsolationForest and LogisticRegression.\n"
        "• Data Wrangling: NumPy 2.5.3 (C-API 2.0 compatible) and Pandas 3.0.5 linked successfully.\n"
        "• Database & API: SQLAlchemy 2.0.52, psycopg2-binary 2.9.13, and FastAPI 0.128.8 pinned in requirements.txt."
    )

    doc.add_paragraph(
        "Repository Architecture & Security Guardrails:\n"
        "To prevent accidental leaks of credentials, multi-gigabyte datasets, or large binary checkpoints, a strict .gitignore "
        "was established, excluding: virtual environments (.venv/), raw and processed datasets (data/raw/*, data/processed/*), "
        "serialized model checkpoints (models/* except metadata JSON), environment secrets (.env), and SQLite/temporary databases. "
        "Clean .gitkeep markers preserve the directory tree across backend/, ml_engine/, database/, iot_firmware/, tests/, scripts/, and docs/."
    )

    # --------------------------------------------------------------------------
    # 4. MILESTONE 1: DATASET ACQUISITION & EDA
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("4. Milestone 1: PaySim Dataset Acquisition & EDA", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "PaySim is a synthetic mobile-money transaction dataset generated using the PaySim multi-agent simulator and based on "
        "aggregated characteristics of real financial activity (Lopez-Rojas et al.). While IEEE-CIS contains 434 heavily anonymized "
        "features with extensive missingness, PaySim provides 11 transparent columns and explicit entity relationships (nameOrig, nameDest) "
        "that directly map into relational database schemas. We emphasize that PaySim is a synthetic simulation; A.R.G.U.S. was not trained "
        "directly on real banking customer records."
    )

    add_callout(
        doc,
        "SYNTHETIC SIMULATOR LIMITATION NOTICE:\n"
        "PaySim is generated via agent-based simulation. Simulated fraud agents follow rigid, deterministic heuristics (specifically: "
        "draining sender accounts completely and often bypassing recipient balances). As established during the Milestone 3A audit, "
        "machine learning models achieve exceptionally high scores on PaySim because these synthetic rules are mathematically distinct. "
        "We explicitly do NOT present these scores as proof of real-world fraud detection performance, where human adversaries continuously "
        "mutate their strategies.",
        "ACADEMIC RIGOR & SIMULATION ACKNOWLEDGMENT"
    )

    # Dataset Profile Table
    prof_table = doc.add_table(rows=12, cols=4)
    p_col_w = [1.8, 1.2, 1.5, 2.0]
    prof_data = [
        ("Column Name", "Data Type", "Missing (%)", "Observed Range / Semantics"),
        ("step", "int32", "0.00%", "1 to 743 (~30.95 days simulation, 1 step = 1 hour)"),
        ("type", "category", "0.00%", "5 types: CASH_IN, CASH_OUT, DEBIT, PAYMENT, TRANSFER"),
        ("amount", "float64", "0.00%", "$0.00 to $92,445,516.64 (Transferred funds)"),
        ("nameOrig", "string", "0.00%", "6,353,307 unique sender accounts (C... prefix)"),
        ("oldbalanceOrg", "float64", "0.00%", "$0.00 to $59,585,040.37 (Sender initial balance)"),
        ("newbalanceOrig", "float64", "0.00%", "$0.00 to $49,585,040.37 (Sender ending balance)"),
        ("nameDest", "string", "0.00%", "2,722,362 unique recipients (C... customer / M... merchant)"),
        ("oldbalanceDest", "float64", "0.00%", "$0.00 to $356,015,889.35 (Recipient initial balance)"),
        ("newbalanceDest", "float64", "0.00%", "$0.00 to $356,179,278.92 (Recipient ending balance)"),
        ("isFraud", "int8", "0.00%", "0 (Legitimate: 6,354,407) or 1 (Fraud: 8,213)"),
        ("isFlaggedFraud", "int8", "0.00%", "0 (6,362,604) or 1 (16 rows flagged > $200,000)"),
    ]
    for idx, r_data in enumerate(prof_data):
        row = prof_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(prof_table, p_col_w)

    doc.add_paragraph()

    h2 = doc.add_heading("Empirical EDA Discoveries", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Execution of scripts/inspect_paysim.py on the raw 470.67 MB CSV file revealed several critical empirical findings:\n"
        "1. Complete Data Integrity: Exactly 0 missing values and 0 duplicate rows across all 6,362,620 transactions.\n"
        "2. Extreme Class Imbalance: 8,213 fraudulent transactions versus 6,354,407 legitimate transactions (0.1291% fraud rate; "
        "773.7:1 imbalance ratio). Accuracy is entirely invalidated as an evaluation metric.\n"
        "3. 100% Fraud Concentration in TRANSFER and CASH_OUT: Fraud occurred exclusively within TRANSFER (4,097 cases; 0.7688% fraud rate) "
        "and CASH_OUT (4,116 cases; 0.1840% fraud rate). Zero fraud occurred in PAYMENT (2.15M tx), CASH_IN (1.40M tx), or DEBIT (41K tx).\n"
        "4. Account Liquidation Pattern: In 97.55% of fraudulent transactions (8,012 / 8,213), the sender account was completely drained "
        "to zero (oldbalanceOrg > 0 and newbalanceOrig == 0), compared to only 23.80% in legitimate transactions.\n"
        "5. Destination Mule Account Anomaly: In 49.63% of fraud cases (4,076 / 8,213), both oldbalanceDest and newbalanceDest remained "
        "exactly 0.00 despite hundreds of thousands of dollars being transferred.\n"
        "6. Diurnal Fraud Surge: Legitimate transactions drop significantly at night (trough of ~1,200 tx/hour at 04:00), whereas fraud "
        "attempts remain continuous (~300–370/hour), driving the early morning fraud rate up to 16.2% – 22.3% (versus 0.08% at midday).\n"
        "7. Investigation of isFlaggedFraud: In the raw dataset, exactly 16 transactions were flagged as isFlaggedFraud==1 based on an "
        "archaic rule (transfer > $200,000). It missed 8,197 fraud cases (99.8% false negative rate) and represents a downstream rule output. "
        "It was permanently excluded to eliminate target leakage."
    )

    # Embed All 5 EDA Figures
    figures_dir = Path("docs/figures")
    
    if (figures_dir / "eda_01_transaction_types.png").is_file():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(figures_dir / "eda_01_transaction_types.png"), width=Inches(5.5))
        p_cap = doc.add_paragraph("Figure 1: Overall PaySim Transaction Volume Distribution across Five Operational Types")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.color.rgb = MUTED

    if (figures_dir / "eda_02_fraud_by_type.png").is_file():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(figures_dir / "eda_02_fraud_by_type.png"), width=Inches(5.5))
        p_cap = doc.add_paragraph("Figure 2: Empirical Fraud Volume & Percentage by Transaction Type (100% Concentration in TRANSFER & CASH_OUT)")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.color.rgb = MUTED

    if (figures_dir / "eda_03_amount_distribution.png").is_file():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(figures_dir / "eda_03_amount_distribution.png"), width=Inches(5.5))
        p_cap = doc.add_paragraph("Figure 3: Logarithmic Monetary Value Distribution: Legitimate Commerce vs Fraudulent Exfiltration")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.color.rgb = MUTED

    if (figures_dir / "eda_04_temporal_trend.png").is_file():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(figures_dir / "eda_04_temporal_trend.png"), width=Inches(5.5))
        p_cap = doc.add_paragraph("Figure 4: Temporal Fraud Rate & Transaction Frequency Over the 743-Hour Simulation Timeline")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.color.rgb = MUTED

    if (figures_dir / "eda_05_balance_drainage.png").is_file():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(figures_dir / "eda_05_balance_drainage.png"), width=Inches(5.2))
        p_cap = doc.add_paragraph("Figure 5: Origin Account Liquidation Behavior: Legitimate (23.8%) vs Fraudulent (97.55%) Transactions")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.color.rgb = MUTED

    # --------------------------------------------------------------------------
    # 5. MILESTONE 2: FEATURE ENGINEERING & DATA SPLITTING
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("5. Milestone 2: Feature Engineering & Data Splitting", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "To transform the 11 raw columns into high-performance numerical inputs without introducing future temporal leakage, "
        "a modular feature pipeline was engineered in ml_engine/features/ and executed via scripts/build_paysim_features.py."
    )

    doc.add_paragraph(
        "The Modeling Subspace Strategy:\n"
        "Because EDA verified that 100% of fraud occurred exclusively in TRANSFER and CASH_OUT, the machine learning modeling "
        "subspace was focused specifically on these two transaction types (2,770,409 rows; 100% of fraud labels). However, "
        "PAYMENT, CASH_IN, and DEBIT are retained conceptually at the system architecture level: they will be handled by a dedicated "
        "low-latency business-rule path in the API gateway. The exact rules and operational policy will be finalized during API integration. [DESIGNED / PLANNED]"
    )

    # Feature Table
    feat_table = doc.add_table(rows=19, cols=3)
    f_col_w = [1.8, 2.2, 2.5]
    feat_data = [
        ("Feature Name", "Mathematical Formulation", "Domain Rationale & Leakage Safeguard"),
        ("amount", "Raw numerical float", "Monetary transaction value requested at ingress. Available at runtime. SAFE."),
        ("oldbalanceOrg", "Raw numerical float", "Pre-transaction liquidity of sender from internal ledger. SAFE."),
        ("newbalanceOrig", "Raw numerical float", "Post-transaction sender balance. Monitored via ledger balance error. SAFE."),
        ("oldbalanceDest", "Raw numerical float", "Pre-transaction recipient balance from ledger. SAFE."),
        ("newbalanceDest", "Raw numerical float", "Post-transaction recipient balance from ledger. SAFE."),
        ("log_amount", "log_e(amount + 1.0)", "Compresses extreme monetary positive skew; stabilizes neural network gradients. SAFE."),
        ("is_transfer", "I(type == 'TRANSFER')", "Binary indicator for inter-account transfer routing. SAFE."),
        ("is_cash_out", "I(type == 'CASH_OUT')", "Binary indicator for ATM/agent cash-out liquidation routing. SAFE."),
        ("hour_of_day", "step mod 24", "Diurnal hour (0 to 23). Captures human circadian activity cycles. SAFE."),
        ("hour_sin", "sin(2 * pi * hour / 24)", "Cyclical circular sine encoding ensuring continuity at midnight (23:00 to 00:00). SAFE."),
        ("hour_cos", "cos(2 * pi * hour / 24)", "Cyclical circular cosine encoding ensuring smooth periodic representation. SAFE."),
        ("is_night_transaction", "I(1 <= hour <= 6)", "Flags the 01:00-06:00 window where fraud concentration surges to 22.3%. SAFE."),
        ("orig_balance_error", "newOrig + amount - oldOrig", "Quantifies deviation from double-entry ledger balance conservation on sender. SAFE."),
        ("orig_drain_ratio", "amount / (oldOrig + 1.0)", "Proportion of sender liquidity liquidated (epsilon prevents zero division). SAFE."),
        ("is_full_liquidation", "I(oldOrig > 0 and newOrig == 0)", "Flags 100% account balance drainage (observed in 97.55% of fraud). SAFE."),
        ("dest_balance_error", "oldDest + amount - newDest", "Quantifies deviation from expected balance addition on recipient ledger. SAFE."),
        ("dest_drain_ratio", "amount / (newDest + 1.0)", "Ratio of transferred amount relative to recipient ending balance. SAFE."),
        ("dest_zero_balance_anomaly", "I(oldDest==0 and newDest==0 and amount>0)", "Mule account marker: recipient balances remain zero despite fund arrival. SAFE."),
    ]
    for idx, r_data in enumerate(feat_data):
        row = feat_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(feat_table, f_col_w)

    doc.add_paragraph()

    h2 = doc.add_heading("Target Isolation & Excluded Features", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "To guarantee zero data leakage, strict programmatic assertions verify that the following fields NEVER enter feature matrix X:\n"
        "• isFraud: Ground truth target; strictly isolated as dependent variable y.\n"
        "• isFlaggedFraud: Excluded completely. It is an artificial simulator rule (> $200,000) that missed 99.8% of fraud and represents a downstream policy output.\n"
        "• nameOrig & nameDest: Raw string identifiers (6.35M unique senders). 99.85% of senders appear exactly once; feeding raw IDs into models causes severe memorization.\n"
        "• step: Monotonically increasing counter. Replaced by cyclical diurnal features to prevent linear split drift in tree models.\n"
        "• Sender Behavioral Velocity: Multi-transaction rolling velocity features (e.g., transactions in past hour) were investigated and "
        "deliberately deferred because 99.85% of senders in PaySim appear only once. Computing historical rolling aggregates on single-occurrence "
        "entities creates meaningless sparse columns; velocity features will be implemented via stateful Redis caching in Milestone 8."
    )

    h2 = doc.add_heading("Chronological Temporal Partitioning", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Random k-fold cross-validation was strictly rejected because it leaks future diurnal patterns into past training. "
        "The modeling subspace (2,770,409 rows) was partitioned chronologically along the simulation timeline:"
    )

    split_table = doc.add_table(rows=5, cols=6)
    s_col_w = [1.2, 1.2, 1.2, 1.0, 1.0, 1.0]
    split_data = [
        ("Partition", "Step Range", "Row Count", "Fraud Count", "Fraud Rate", "Imbalance"),
        ("Train", "Steps 1 to 520 (70.0%)", "2,653,729", "5,781", "0.2178%", "458.0 : 1"),
        ("Validation", "Steps 521 to 631 (14.9%)", "78,701", "1,180", "1.4993%", "65.7 : 1"),
        ("Test (Holdout)", "Steps 632 to 743 (15.1%)", "37,979", "1,252", "3.2966%", "29.3 : 1"),
        ("Total Subspace", "Steps 1 to 743 (100%)", "2,770,409", "8,213", "0.2965%", "336.3 : 1"),
    ]
    for idx, r_data in enumerate(split_data):
        row = split_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(split_table, s_col_w)

    doc.add_paragraph()
    doc.add_paragraph(
        "Scaler Isolation Guarantee: Scikit-learn's StandardScaler was fitted strictly on the Train partition (X_train) and "
        "serialized to models/feature_scaler.joblib. The validation and test partitions were transformed using training parameters "
        "without updating mean or variance parameters. Five unit tests in tests/test_features.py pass cleanly."
    )

    # --------------------------------------------------------------------------
    # 6. MILESTONE 3: MODEL PREPARATION & TRAINING PIPELINE
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("6. Milestone 3: Model Development & Training Results", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "A robust fraud detection system cannot rely on a single model architecture. A.R.G.U.S. deploys a four-model strategy "
        "combining supervised linear classification, gradient-boosted decision trees, tree-based unsupervised anomaly detection, "
        "and deep neural reconstruction error:"
    )

    doc.add_paragraph(
        "Model Roles & Behavioral Niches:\n"
        "1. Logistic Regression: Baseline linear fraud classifier. Uses class_weight='balanced' to handle class skew. "
        "Provides an interpretable linear benchmark. Formal probability calibration such as Platt scaling or isotonic regression has not been implemented; "
        "its outputs represent model-estimated fraud probabilities.\n"
        "2. XGBoost: Primary supervised fraud classifier. Captures complex nonlinear feature interactions across balance discrepancies. "
        "Configured with scale_pos_weight = 458.04 (the exact negative-to-positive ratio in training data).\n"
        "3. Isolation Forest: Tree-based unsupervised anomaly detection. Isolates outliers through recursive random feature partitioning; "
        "anomalies tend to be isolated through shorter tree paths without requiring fraud labels, allowing it to detect novel fraud topologies.\n"
        "4. Deep Autoencoder (PyTorch): Deep learning representation model. Comprises a symmetric bottleneck architecture:\n"
        "   Input(18) → Dense(64) → Dense(32) → Latent Bottleneck(16) → Dense(32) → Dense(64) → Output(18)\n"
        "   Trained strictly on legitimate transactions with Mean Squared Error (MSE) reconstruction loss. The 16-dimensional bottleneck "
        "   serves as the compact latent space representation of normal transaction dynamics. Transactions producing elevated reconstruction MSE "
        "   represent anomalous behavior deviating from the normal manifold."
    )

    h2 = doc.add_heading("Empirical Evaluation Metrics (Validation & Test Sets)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "The following results represent verified empirical metrics logged during Milestone 3 and audited in Milestone 3A. "
        "Optimal decision thresholds were selected on the Validation partition to maximize the F1-score and evaluated on the Test holdout:"
    )

    # Performance Table - Validation Set
    doc.add_paragraph("Table A: Model Performance on Validation Partition (78,701 transactions; 1,180 fraud cases)")
    val_table = doc.add_table(rows=5, cols=7)
    m_col_w = [1.6, 0.8, 0.8, 0.8, 0.8, 0.8, 0.9]
    val_data = [
        ("Model Architecture", "PR-AUC", "ROC-AUC", "Precision", "Recall", "F1-Score", "Opt. Thresh"),
        ("Logistic Regression", "0.9400", "0.9976", "0.9753", "0.8364", "0.9005", "0.9902"),
        ("XGBoost Classifier", "0.9997", "1.0000", "1.0000", "0.9992", "0.9996", "0.9425"),
        ("Isolation Forest", "0.5943", "0.9248", "0.6014", "0.6381", "0.6192", "0.3828"),
        ("Deep Autoencoder", "0.7402", "0.9440", "0.9199", "0.6619", "0.7698", "0.9999"),
    ]
    for idx, r_data in enumerate(val_data):
        row = val_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(val_table, m_col_w)

    doc.add_paragraph()

    # Performance Table - Test Set
    doc.add_paragraph("Table B: Model Performance on Test Holdout Partition (37,979 transactions; 1,252 fraud cases)")
    test_table = doc.add_table(rows=5, cols=7)
    test_data = [
        ("Model Architecture", "PR-AUC", "ROC-AUC", "Precision", "Recall", "F1-Score", "FPR"),
        ("Logistic Regression", "0.9512", "0.9946", "0.9826", "0.8099", "0.8879", "0.0005"),
        ("XGBoost Classifier", "1.0000", "1.0000", "1.0000", "0.9992", "0.9996", "0.0000"),
        ("Isolation Forest", "0.7329", "0.9163", "0.8576", "0.6254", "0.7233", "0.0035"),
        ("Deep Autoencoder", "0.7835", "0.9302", "0.9763", "0.6573", "0.7857", "0.0005"),
    ]
    for idx, r_data in enumerate(test_data):
        row = test_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(test_table, m_col_w)

    doc.add_paragraph()

    doc.add_paragraph(
        "Important Contextualization on XGBoost Performance:\n"
        "XGBoost achieved near-perfect performance on the PaySim test partition (Test PR-AUC 1.0000, F1 0.9996). "
        "This result is strongly influenced by deterministic transaction-generation patterns in the synthetic PaySim environment "
        "and should not be generalized directly to real-world financial fraud. In live production banking, fraud strategies mutate, "
        "attack vectors evolve, and balance behaviors do not follow rigid deterministic formulas."
    )

    doc.add_paragraph(
        "Deep Autoencoder Architecture & Training Specifications:\n"
        "• Network Topology: Symmetric deep bottleneck:\n"
        "  Input(18) → Linear(18, 64) + BatchNorm1d + LeakyReLU(0.1)\n"
        "  → Linear(64, 32) + BatchNorm1d + LeakyReLU(0.1)\n"
        "  → Linear(32, 16) [16-dimensional Latent Bottleneck]\n"
        "  → Linear(16, 32) + BatchNorm1d + LeakyReLU(0.1)\n"
        "  → Linear(32, 64) + BatchNorm1d + LeakyReLU(0.1)\n"
        "  → Linear(64, 18) [Reconstruction Output]\n"
        "• Loss Function: Mean Squared Error (MSE) between input and reconstructed feature vector.\n"
        "• Training Regime: Trained on CPU using 2,647,948 legitimate transactions from the Train partition.\n"
        "• Optimizer & Schedule: Adam optimizer (lr=0.001), batch size 2,048.\n"
        "• Convergence: 15 epochs attempted; early stopping triggered at Epoch 11 with best validation normal MSE of 0.002083.\n"
        "• Error Distribution: Normal validation reconstruction error exhibited mean = 0.002975, std = 0.014327, and P95 = 0.009046."
    )

    # --------------------------------------------------------------------------
    # 7. MILESTONE 3A: MODEL AUDIT & FORENSIC INVESTIGATION
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("7. Milestone 3A: Model Result & Leakage Audit", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "In response to the near-perfect metrics achieved by XGBoost (Test PR-AUC 1.0000, F1 0.9996; TP=1,251, FP=0, TN=36,727, FN=1), "
        "a comprehensive forensic audit was conducted in Milestone 3A. In machine learning, near-perfect scores frequently indicate "
        "target leakage, future information contamination, or evaluation bugs. Milestone 3A audited the source code, feature "
        "distributions, and booster split gains to establish technical validity."
    )

    h2 = doc.add_heading("Feature Importance & Gain Concentration", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "An inspection of the XGBoost booster gain metrics (the relative contribution of each feature to decision tree splits) "
        "uncovered an extreme concentration of predictive power:"
    )

    gain_table = doc.add_table(rows=6, cols=3)
    g_col_w = [2.0, 1.8, 2.7]
    gain_data = [
        ("Feature Name", "Gain Contribution (%)", "Behavioral / Synthetic Rationale"),
        ("orig_balance_error", "66.71%", "Double-entry violation: amount exceeds old balance or doesn't decrement cleanly."),
        ("newbalanceOrig", "11.07%", "Zero-balance endpoint: 97.55% of simulated frauds drain sender to exactly 0.00."),
        ("orig_drain_ratio", "9.01%", "Liquidation proportion: fraud transactions overwhelmingly exhibit drain_ratio == 1.0."),
        ("is_full_liquidation", "5.47%", "Direct indicator for complete balance drainage."),
        ("Top 4 Combined", "92.25%", "Four sender balance features account for over 92% of the booster's total split gain."),
    ]
    for idx, r_data in enumerate(gain_data):
        row = gain_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(gain_table, g_col_w)

    doc.add_paragraph()

    doc.add_paragraph(
        "Why is XGBoost Performance So High? The Synthetic Simulator Reality:\n"
        "The forensic audit conclusively verified that there is ZERO TARGET LEAKAGE and ZERO TEMPORAL LEAKAGE. All 18 features "
        "are strictly derivable from the ingress transaction payload and current ledger balances at transaction time. The extraordinary "
        "performance is an inherent artifact of the PaySim multi-agent simulation: the synthetic fraud agents were programmed with rigid, "
        "deterministic heuristics (specifically: always empty the victim account completely, and frequently discard destination ledger balances). "
        "Because these heuristic rules are deterministic, gradient-boosted trees easily partition the feature space with near-zero error. "
        "In a real-world banking environment, human fraudsters continuously vary transaction amounts, leave residual balances, and utilize "
        "complex money-mule layering to blend with normal commerce. Thus, the near-perfect metrics must be understood as an artifact of "
        "the simulation rather than an infallible real-world fraud detector."
    )

    h2 = doc.add_heading("Forensic Analysis of the Singular False Negative (Test Row 3583)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Across all 37,979 test transactions, XGBoost produced exactly ONE false negative (Row index 3583). A forensic autopsy "
        "of this transaction demonstrates why a multi-model ensemble is critical:"
    )

    fn_table = doc.add_table(rows=8, cols=2)
    fn_col_w = [2.2, 4.3]
    fn_data = [
        ("Transaction Attribute", "Observed Empirical Value"),
        ("Transaction Type & Step", "TRANSFER at Step 646 (Hour 22:00 Diurnal)"),
        ("Transaction Amount", "$399,045.08"),
        ("Sender Balances", "oldbalanceOrg = $10,399,045.08 | newbalanceOrig = $10,399,045.08 (Balance not debited)"),
        ("Recipient Balances", "oldbalanceDest = $0.00 | newbalanceDest = $0.00 (Mule zero-balance anomaly)"),
        ("Drain Ratio & Balance Error", "orig_drain_ratio = 0.0384 (3.84% of funds) | orig_balance_error = $399,045.08"),
        ("Individual Model Inferences", "XGBoost = 0.7654 | Autoencoder = 1.0000 (MSE: 44.84) | Isolation Forest = 0.7394 | Logistic Regression = 1.0000"),
        ("Risk Engine Evaluation", "Aggregated Risk Score = 84.36 | Policy Action = BLOCK (Threshold >= 70.0)"),
    ]
    for idx, (label, val) in enumerate(fn_data):
        row = fn_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(fn_col_w[0])
        c1.width = Inches(fn_col_w[1])
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "F7FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(9.0)
        r0.font.color.rgb = NAVY
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.0)
        r1.font.color.rgb = CHARCOAL

    doc.add_paragraph()
    doc.add_paragraph(
        "Forensic Insight: In this transaction, the fraudster did NOT drain the account to zero; they attempted a transfer of $399,045.08 "
        "from a massive $10.4M account, leaving the balance untouched in the simulation (oldbalanceOrg == newbalanceOrig == $10,399,045.08). "
        "Because XGBoost learned that PaySim fraud almost universally exhibits full liquidation (92.25% gain concentration), its model-estimated "
        "probability dropped to 0.7654. Under XGBoost's standalone binary threshold of 0.9425, it was misclassified as legitimate (FN). "
        "However, under the A.R.G.U.S. multi-signal Risk Engine, the combination of XGBoost (0.7654), Autoencoder reconstruction anomaly detection "
        "(1.0000, MSE: 44.84), Isolation Forest (0.7394), and Logistic Regression (1.0000) yields an exact aggregate Risk Score of 84.36. "
        "This places the transaction decisively into the BLOCK tier (Score >= 70.0), completely preventing the fraud from escaping undetected!"
    )

    h2 = doc.add_heading("PyTorch Tensor Writability Bug Fix", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "During the audit, a UserWarning was identified in ml_engine/models/autoencoder.py: 'The given NumPy array is not writable, "
        "and PyTorch does not support non-writable tensors'. This occurred because Scikit-learn feature slices generated non-contiguous, "
        "read-only memory views before conversion to PyTorch tensors. The issue was resolved cleanly without retraining by modifying "
        "predict_reconstruction_error() to create an explicit writable copy via torch.from_numpy(np.array(x_scaled, copy=True)). "
        "Subsequent test execution verified 12/12 unit tests passing with zero warnings."
    )

    # --------------------------------------------------------------------------
    # 8. MILESTONE 4: RISK SCORING & DECISION ENGINE
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("8. Milestone 4: Risk Scoring & Decision Engine", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "Milestone 4 implemented the central intelligence hub of A.R.G.U.S.: the Risk Scoring & Decision Engine (ml_engine/risk/). "
        "The Risk Engine provides a multi-signal operational framework that combines supervised classification and independent anomaly "
        "signals into a continuous risk score and tri-state decision. It normalizes heterogeneous signals across all four models into "
        "a bounded continuous scale [0.0, 1.0], combines them via an empirically weighted sum, scales the result to a 0–100 Risk Score, "
        "and triggers an operational tri-state decision policy."
    )

    doc.add_paragraph(
        "Signal Normalization Mechanics:\n"
        "• Supervised Signals (XGBoost & Logistic Regression): Outputs represent model-estimated fraud probabilities residing in [0.0, 1.0]. "
        "Formal probability calibration such as Platt scaling or isotonic regression has not been implemented.\n"
        "• Isolation Forest Signal: Scikit-learn's score_samples() outputs tree-based isolation scores where smaller values indicate outliers. "
        "The signal is normalized using empirical training distribution bounds [0.3395, 0.7949]:\n"
        "  Normalized IF Signal = clip((score - min_bound) / (max_bound - min_bound), 0.0, 1.0)\n"
        "• Deep Autoencoder Signal: Raw Mean Squared Error (MSE) reconstruction loss is mapped to [0.0, 1.0] using validation normal statistics "
        "(mean = 0.002975, std = 0.014327, P95 = 0.009046):\n"
        "  Normalized AE Signal = clip(0.5 + 0.5 * (mse - val_mean) / (3.0 * val_std), 0.0, 1.0)"
    )

    doc.add_paragraph(
        "Ensemble Aggregation Weights & Risk Scoring Formula:\n"
        "The individual normalized signals S_i are combined using conservative ensemble weights reflecting each model's audited reliability:\n"
        "• XGBoost (w_xgb = 0.50): Dominant supervised fraud classifier.\n"
        "• Deep Autoencoder (w_ae = 0.20): Deep representation anomaly detector (18 → 64 → 32 → 16 → 32 → 64 → 18).\n"
        "• Isolation Forest (w_if = 0.15): Tree-based unsupervised anomaly detector.\n"
        "• Logistic Regression (w_lr = 0.15): Baseline linear fraud classifier.\n\n"
        "Mathematical Formulation:\n"
        "  Aggregated Risk Signal = 0.50 * S_xgb + 0.20 * S_ae + 0.15 * S_if + 0.15 * S_lr\n"
        "  Risk Score = Aggregated Risk Signal * 100.0  ∈ [0.00, 100.00]"
    )

    doc.add_paragraph(
        "Tri-State Operational Decision Policies:\n"
        "• Score < 40.0 → APPROVE: Low risk. Immediate settlement with near-zero latency.\n"
        "• 40.0 <= Score < 70.0 → REVIEW: Elevated risk. Routed to fraud analysts with observable risk factors and model explanations.\n"
        "• Score >= 70.0 → BLOCK: Critical risk. Transaction rejected by the current decision policy. (Account freezing or forensic hold "
        "is noted strictly as a potential future operational policy)."
    )

    h2 = doc.add_heading("Empirical Evaluation on the PaySim Validation Partition (78,701 Transactions)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "The Risk Engine was evaluated across the entire Validation partition (78,701 transactions; 1,180 ground-truth fraud cases). "
        "Validation-set operational evaluation demonstrates strong triage performance on the simulated data:"
    )

    dec_table = doc.add_table(rows=4, cols=6)
    d_col_w = [1.2, 1.1, 1.1, 1.1, 1.0, 1.0]
    dec_data = [
        ("Decision", "Total Volume", "% of Traffic", "Legitimate", "Fraud Caught", "Fraud Hit Rate"),
        ("APPROVE", "77,473", "98.44%", "77,472", "1 (FN)", "0.00%"),
        ("REVIEW", "165", "0.21%", "49", "116", "70.30%"),
        ("BLOCK", "1,063", "1.35%", "0 (0 FP)", "1,063", "100.00%"),
    ]
    for idx, r_data in enumerate(dec_data):
        row = dec_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(dec_table, d_col_w)

    doc.add_paragraph()

    doc.add_paragraph(
        "Operational Evaluation Analysis:\n"
        "• 99.92% Total Fraud Capture: 1,179 out of 1,180 fraud cases were intercepted (1,063 blocked immediately + 116 routed to review).\n"
        "• Zero False Rejections: Exactly 0 legitimate transactions were rejected in the BLOCK tier (0.00% false positive block rate on validation data).\n"
        "• Low Human Review Burden: Only 165 transactions (0.21% of total traffic) require manual investigation, with a 70.30% fraud density in the queue.\n"
        "• Multi-Signal Framework: Rather than claiming the ensemble 'dramatically outperforms' XGBoost (which already performs exceptionally "
        "strongly on PaySim), the Risk Engine's primary value is providing an operational decision framework that balances automated blockage "
        "with human-in-the-loop review for borderline transactions."
    )

    h2 = doc.add_heading("Explainability & Domain Entity Architecture", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Every evaluation produces an immutable RiskAssessment domain entity containing:\n"
        "• Individual Model Signals: Raw scores and normalized signals for XGBoost, Autoencoder, Isolation Forest, and Logistic Regression.\n"
        "• Observable Feature Patterns: Human-readable diagnostic indicators ('Account liquidation pattern: 100% drained', 'Mule destination balance anomaly', 'Night window 01:00-06:00', 'Ledger balance conservation violation').\n"
        "• System Metadata: engine_version ('1.0.0'), evaluated_by ('ml_ensemble' or 'business_rule_fast_path'), timestamp, and tracking ID."
    )

    # --------------------------------------------------------------------------
    # 9. PLANNED ARCHITECTURES: DBMS, SECURITY, IOT & OOPS
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("9. Planned System Architectures (Milestones 5–7)", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "The following architectural designs have been conceptualized and analyzed during Milestones 0–4 but are explicitly "
        "PLANNED FOR FUTURE IMPLEMENTATION. Neither database migrations, live REST endpoints, nor physical IoT flashing have occurred."
    )

    h2 = doc.add_heading("Planned DBMS Architecture — Milestone 5 (PLANNED — NOT YET IMPLEMENTED)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Conceptual PostgreSQL relational architecture designed for Milestone 5. The schema follows Third Normal Form (3NF) "
        "to ensure ACID-compliant transaction persistence and auditability across 11 core entities:"
    )

    # DBMS Entities Table
    db_table = doc.add_table(rows=12, cols=3)
    db_col_w = [1.8, 1.8, 2.9]
    db_data = [
        ("Entity / Table", "Primary Key & Foreign Keys", "Domain Purpose & Relational Attributes"),
        ("users", "user_id (UUID PK)", "User credentials, status, role_id (FK -> roles), created_at."),
        ("merchants", "merchant_id (UUID PK)", "Merchant business profile, category code, risk tier, registered_at."),
        ("devices", "device_id (UUID PK)", "POS/IoT hardware registry, merchant_id (FK), public_key, status, firmware_ver."),
        ("transactions", "tx_id (UUID PK)", "Monetary ledger: user_id (FK), merchant_id (FK), amount, type, step, created_at."),
        ("transaction_features", "feature_id (UUID PK)", "18 engineered features linked to tx_id (FK -> transactions)."),
        ("risk_assessments", "assessment_id (UUID PK)", "Aggregated risk score (0-100), decision, tx_id (FK), evaluated_at."),
        ("alerts", "alert_id (UUID PK)", "High-risk alerts: assessment_id (FK), severity, status, assigned_analyst."),
        ("decisions", "decision_id (UUID PK)", "Operational action: assessment_id (FK), policy (APPROVE/REVIEW/BLOCK), reason."),
        ("user_behavior_profiles", "profile_id (UUID PK)", "Stateful behavioral baselines: user_id (FK), avg_amount, tx_frequency."),
        ("audit_logs", "log_id (BIGSERIAL PK)", "Immutable audit trail: entity_type, action, actor_id, prev_hash, current_hash."),
        ("model_versions", "version_id (UUID PK)", "Model metadata registry: model_name, training_date, validation_pr_auc, active."),
    ]
    for idx, r_data in enumerate(db_data):
        row = db_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(db_table, db_col_w)

    doc.add_paragraph()

    h2 = doc.add_heading("Planned Information Security Architecture (DESIGNED / PLANNED)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Security mechanisms designed for future implementation across Milestones 5–7:\n"
        "• Authentication & Authorization: OAuth2 with short-lived JWT bearer tokens (15-minute expiration) and secure refresh tokens. [PLANNED / M6]\n"
        "• Role-Based Access Control (RBAC): Strict permission boundaries across Admin (system configuration), Analyst (review queue triage), and Auditor (read-only audit log inspection). [PLANNED / M6]\n"
        "• API Sanitization: Strict Pydantic input schemas, boundary validation, and rate limiting (token bucket algorithm). [PLANNED / M6]\n"
        "• Password Security: Adaptive password hashing via Argon2id with work factor calibration. [PLANNED / M5]\n"
        "• Cryptographic IoT Integrity: HMAC-SHA256 signature calculated over device_id + timestamp + amount using secure per-device pre-shared keys. [PLANNED / M7]\n"
        "• Secure Communication: Mandatory Transport Layer Security (TLS 1.3) with forward secrecy. [PLANNED / M6]\n"
        "• Tamper-Evident Audit Logging: Cryptographic forward-chained SHA-256 hash chains across audit_logs to guarantee tamper detection. [PLANNED / M5]"
    )

    h2 = doc.add_heading("Planned IoT Architecture — Milestone 7 (PLANNED — PENDING)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "Physical ESP32 integration is pending. The planned hardware architecture comprises:\n"
        "• Microcontroller: Espressif ESP32-WROOM-32 (Dual-core 240 MHz, 520 KB SRAM, hardware cryptographic accelerator).\n"
        "• Physical Tamper Detection: Normally-closed hardware micro-switch connected to GPIO with pull-up resistor; triggering sends immediate tamper_status = 1 telemetry.\n"
        "• Telemetry Payload: device_id, merchant_id, hardware RTC timestamp, device_status, tamper_status, firmware_version, and Wi-Fi RSSI signal strength.\n"
        "• Cryptographic Signing: Internal HMAC-SHA256 calculation offloaded to ESP32 hardware cryptographic engine."
    )

    h2 = doc.add_heading("Object-Oriented Programming (OOP) Architecture", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "The project applies object-oriented principles across application domain models and service layers:\n"
        "• Domain Object Models: Encapsulated entities including User, Transaction, Device, Merchant, RiskAssessment, Alert, Decision, and AuditLog.\n"
        "• Service Abstractions: Decoupled service layer including TransactionService (ingress and validation), RiskAssessmentService (orchestrates scoring), FraudDetectionService (wraps model inference), AuthenticationService (tokens and credentials), and AuditService (immutable logging).\n"
        "• Implemented vs Planned OOP Components: Core ML modules (TransactionRiskEngine, ModelSignalNormalizer, RiskAssessment, DeepAutoencoder, RiskEngineConfig) are fully implemented; backend service abstractions will be realized in FastAPI during Milestone 6."
    )

    # --------------------------------------------------------------------------
    # 10. CURRENT STATUS, LIMITATIONS & NEXT MILESTONES
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("10. Status Summary, Limitations & Next Milestones", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "Comprehensive Project Status Audit (Implemented vs Designed vs Planned):"
    )

    # Status Table
    status_table = doc.add_table(rows=12, cols=3)
    s_widths = [1.8, 1.4, 3.3]
    status_data = [
        ("Component / Milestone", "Implementation Status", "Technical Verification & Verification Details"),
        ("M0 — Environment Setup", "IMPLEMENTED", "Python 3.13.7 virtual environment, requirements pinned, Git initialized at commit 0f1b9c7."),
        ("M1 — PaySim EDA", "IMPLEMENTED", "PaySim 6.36M rows verified, 0 missing, 100% fraud in TRANSFER/CASH_OUT, 5 figures generated."),
        ("M2 — Feature Engineering & Split", "IMPLEMENTED", "18-feature pipeline, chronological split (1-520/521-631/632-743), scaler isolated to Train."),
        ("M3 — ML/DL Model Training", "IMPLEMENTED", "XGBoost, Deep Autoencoder (18-64-32-16-32-64-18), Isolation Forest, Logistic Regression trained."),
        ("M3A — Model Audit & Bug Fix", "IMPLEMENTED", "Forensic audit completed, 92.25% gain explained, PyTorch tensor writability bug fixed."),
        ("M4 — Risk Scoring & Decision", "IMPLEMENTED", "Weighted aggregation (0.50/0.20/0.15/0.15), 0-100 score, tri-state policy, 20/20 unit tests."),
        ("PostgreSQL Database Architecture", "DESIGNED / ANALYZED", "Conceptual 3NF relational schema designed across 11 tables for Milestone 5."),
        ("FastAPI Backend Architecture", "DESIGNED / ANALYZED", "REST API architecture designed for Milestone 6; endpoints not yet implemented."),
        ("ESP32 IoT POS Architecture", "DESIGNED / ANALYZED", "Firmware, tamper switch, and HMAC protocol designed for Milestone 7; physical hardware pending."),
        ("Information Security Architecture", "DESIGNED / ANALYZED", "JWT, RBAC, TLS, and audit hashing designed; code-level environment hygiene implemented."),
        ("OOP Domain Architecture", "DESIGNED / ANALYZED", "Domain entities and service abstractions designed; ML engine classes implemented."),
    ]
    for idx, r_data in enumerate(status_data):
        row = status_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(status_table, s_widths)

    doc.add_paragraph()

    h2 = doc.add_heading("Current System Limitations", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "1. Synthetic Nature of PaySim: Fraud agents follow rigid, deterministic liquidation rules. Real-world adversaries vary amounts and strategies.\n"
        "2. Absence of Multi-Transaction Velocity in Static CSV: 99.85% of PaySim senders appear only once. Cross-session velocity tracking "
        "requires stateful Redis caching to be implemented in Milestone 8.\n"
        "3. IoT Telemetry Disconnection: Physical ESP32 hardware and live telemetry are not yet integrated into the scoring pipeline.\n"
        "4. Static Ensemble Weights: Ensemble weights (0.50 / 0.20 / 0.15 / 0.15) represent initial baseline configurations rather than dynamically optimized hyper-parameters.\n"
        "5. Uncalibrated Probabilities: Supervised model outputs represent raw model-estimated probabilities; formal probability calibration has not yet been fitted."
    )

    h2 = doc.add_heading("Immediate Next Milestones", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = BLUE

    doc.add_paragraph(
        "• Milestone 5: PostgreSQL Database Integration — Deploy local PostgreSQL database, execute SQLAlchemy 3NF migrations, create users, "
        "accounts, merchants, transactions, risk_assessments, and immutable audit_logs tables.\n"
        "• Milestone 6: FastAPI Application Gateway — Build modular asynchronous REST endpoints (/api/v1/assess, /api/v1/review, /api/v1/auth) "
        "with JWT authentication, role-based access control (Admin, Analyst, Auditor), and input validation.\n"
        "• Milestone 7: ESP32 IoT POS Terminal Integration — Develop C++/MicroPython firmware on ESP32 microcontroller with hardware tamper switch, "
        "cryptographic HMAC payload signing, and secure TLS transmission to the FastAPI gateway.\n"
        "• Milestone 8: Frontend Dashboard — Responsive analyst portal with real-time risk feed, queue review, and explainability heatmaps.\n"
        "• Milestone 9: Continuous Learning Pipeline — Automated drift detection and periodic retraining pipeline."
    )

    # --------------------------------------------------------------------------
    # 11. GIT COMMIT HISTORY AUDIT
    # --------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading("11. Version Control Audit & Verification", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = NAVY

    doc.add_paragraph(
        "All development progress is fully tracked in Git with semantic commit messages. "
        "Raw datasets and model binaries remain strictly gitignored to preserve clean version control:"
    )

    git_table = doc.add_table(rows=7, cols=3)
    g_col_w = [1.2, 2.5, 2.8]
    git_data = [
        ("Commit Hash", "Commit Message", "Key Artifacts Delivered"),
        ("0f1b9c7", "chore: initialize ARGUS development environment", "Project initialization, virtual environment, .gitignore, README.md, requirements.txt"),
        ("9f52b9d", "feat: acquire and profile PaySim dataset", "PaySim EDA script, profile report, 5 diagnostic figures, dataset summary JSON"),
        ("b7672c0", "feat: engineer PaySim features and temporal splits", "18-feature pipeline, temporal split (train/val/test), StandardScaler fitted, 5 unit tests"),
        ("f0739b0", "feat: prepare fraud detection model training pipeline", "Logistic, XGBoost, Isolation Forest, Autoencoder model wrappers & training CLI scripts"),
        ("37891a3", "audit: validate model results and feature leakage", "Forensic audit report, false-negative inspection, PyTorch array writability fix"),
        ("55f1094", "feat: implement ARGUS risk scoring and decision engine", "Risk engine package (config, signals, aggregator, decision, explanations), 20 unit tests"),
    ]
    for idx, r_data in enumerate(git_data):
        row = git_table.rows[idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].text = val
    style_table(git_table, g_col_w)

    doc.add_paragraph()
    doc.add_paragraph(
        "Unit Test Verification Status: All 20 unit tests pass cleanly in 8.91 seconds via pytest. Git working tree is clean."
    )

    # Save documents
    output_path_revised = Path("docs/ARGUS_Comprehensive_Technical_Report_REVISED.docx")
    output_path_standard = Path("docs/ARGUS_Comprehensive_Technical_Report.docx")
    saved_revised = False
    try:
        doc.save(output_path_revised)
        print(f"[+] Revised report successfully generated at: {output_path_revised}")
        saved_revised = True
    except PermissionError:
        alt_path = Path("docs/ARGUS_Comprehensive_Technical_Report_REVISED_UPDATED.docx")
        doc.save(alt_path)
        print(f"[!] {output_path_revised} is currently open in Microsoft Word. Saved updated report to: {alt_path}")

    try:
        doc.save(output_path_standard)
        print(f"[+] Standard report updated at: {output_path_standard}")
    except PermissionError:
        print(f"[!] {output_path_standard} is also locked by Word.")
        
    return output_path_revised if saved_revised else Path("docs/ARGUS_Comprehensive_Technical_Report_REVISED_UPDATED.docx")


if __name__ == "__main__":
    build_report()
